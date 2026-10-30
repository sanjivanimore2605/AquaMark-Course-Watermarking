# Complete Academic Project Report DOCX Generator for AquaMark
import os
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from report_data_front import FRONT_MATTER, ABSTRACT_TEXT
import report_data_ch1_4 as ch1_4
import report_data_ch5_8 as ch5_8
import report_data_ch9_13 as ch9_13
import report_data_appendices as appx

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def generate_academic_report_docx(output_path):
    doc = Document()

    # Set 1-inch margins (standard university format)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Styles
    primary_color = RGBColor(15, 23, 42)    # Slate 900
    accent_color = RGBColor(13, 148, 136)   # Teal 600
    body_color = RGBColor(30, 41, 59)      # Slate 800

    def add_p(text, style='Normal', space_before=0, space_after=6, line_spacing=1.5, bold=False, italic=False, color=body_color, align=WD_ALIGN_PARAGRAPH.JUSTIFY, font_size=12):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_spacing
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.color.rgb = color
        return p

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = primary_color
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = accent_color
        return p

    def add_bullet(text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.25
        run = p.add_run(text)
        run.font.name = "Times New Roman"
        run.font.size = Pt(11)
        run.font.color.rgb = body_color
        return p

    def add_table_grid(header_cols, rows_data):
        tbl = doc.add_table(rows=len(rows_data) + 1, cols=len(header_cols))
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        # Header
        for c_idx, h_text in enumerate(header_cols):
            cell = tbl.cell(0, c_idx)
            set_cell_background(cell, "0F172A")
            set_cell_margins(cell, 80, 80, 100, 100)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(h_text)
            r.font.name = "Times New Roman"
            r.font.bold = True
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(255, 255, 255)
        # Rows
        for r_idx, r_data in enumerate(rows_data, start=1):
            bg = "FFFFFF" if r_idx % 2 != 0 else "F8FAFC"
            for c_idx, val in enumerate(r_data):
                cell = tbl.cell(r_idx, c_idx)
                set_cell_background(cell, bg)
                set_cell_margins(cell, 60, 60, 90, 90)
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                r = p.add_run(str(val))
                r.font.name = "Times New Roman"
                r.font.size = Pt(9)
                r.font.color.rgb = body_color
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    def add_code_block(code_str):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        set_cell_background(cell, "0F172A")
        set_cell_margins(cell, 90, 90, 120, 120)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(code_str)
        r.font.name = "Courier New"
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(226, 232, 240)
        doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # ==================== COVER PAGE ====================
    add_p(FRONT_MATTER["college_name"].upper(), space_before=20, space_after=4, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=14)
    add_p(FRONT_MATTER["department"].upper(), space_after=24, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=12)
    
    add_p("A PROJECT REPORT ON", space_after=12, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=12)
    add_p(FRONT_MATTER["project_title"].upper(), space_after=8, bold=True, color=primary_color, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=18)
    add_p(f"Topic: {FRONT_MATTER['official_topic']}", space_after=36, italic=True, color=accent_color, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=13)

    add_p("Submitted in partial fulfillment of the requirements\nfor the award of the degree of", space_after=10, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=11)
    add_p("BACHELOR OF TECHNOLOGY", space_after=4, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=14)
    add_p("IN\nCOMPUTER SCIENCE AND ENGINEERING (CYBER SECURITY)", space_after=36, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=12)

    add_p(f"Submitted by:\n{FRONT_MATTER['student_name']}\nPRN: {FRONT_MATTER['roll_number']}\n{FRONT_MATTER['year_semester']}", space_after=24, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=12)
    add_p(f"Under the Guidance of:\n{FRONT_MATTER['guide_name']}", space_after=36, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=12)
    add_p(f"{FRONT_MATTER['university_name']}\nAcademic Year {FRONT_MATTER['academic_year']}", space_after=0, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=12)

    doc.add_page_break()

    # ==================== CERTIFICATE PAGE ====================
    add_p(FRONT_MATTER["college_name"].upper(), space_before=10, space_after=4, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=14)
    add_p(FRONT_MATTER["department"], space_after=24, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=12)
    add_p("BONAFIDE CERTIFICATE", space_after=20, bold=True, color=primary_color, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=16)

    cert_text = (
        f"This is to certify that this project report entitled \"{FRONT_MATTER['project_title']}\" is the bonafide work "
        f"carried out by {FRONT_MATTER['student_name']} (PRN: {FRONT_MATTER['roll_number']}) in partial fulfillment "
        f"of the requirements for the award of the degree of Bachelor of Technology in Computer Science and Engineering "
        f"(Cyber Security) of {FRONT_MATTER['university_name']} during the academic year {FRONT_MATTER['academic_year']}.\n\n"
        "The project work has been carried out under my supervision and guidance. The contents of this report have not been "
        "submitted to any other University or Institute for the award of any degree or diploma."
    )
    add_p(cert_text, space_after=40, line_spacing=1.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    # Guide & HOD Signatures
    tbl_cert = doc.add_table(rows=1, cols=2)
    tbl_cert.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell_guide = tbl_cert.cell(0, 0)
    cell_hod = tbl_cert.cell(0, 1)
    
    p_g = cell_guide.paragraphs[0]
    p_g.paragraph_format.line_spacing = 1.15
    p_g.add_run("_________________________\n").bold = True
    p_g.add_run(f"Internal Guide:\n{FRONT_MATTER['guide_name']}\n{FRONT_MATTER['department']}")
    
    p_h = cell_hod.paragraphs[0]
    p_h.paragraph_format.line_spacing = 1.15
    p_h.add_run("_________________________\n").bold = True
    p_h.add_run(f"Head of Department:\n{FRONT_MATTER['head_of_department']}")

    add_p("\nSubmitted for the University Project Examination held on: ___________________", italic=True, align=WD_ALIGN_PARAGRAPH.LEFT)
    
    tbl_exam = doc.add_table(rows=1, cols=2)
    tbl_exam.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_in = tbl_exam.cell(0, 0)
    c_ex = tbl_exam.cell(0, 1)
    p_in = c_in.paragraphs[0]
    p_in.add_run("\n\n_________________________\nInternal Examiner\nSignature & Date").bold = True
    p_ex = c_ex.paragraphs[0]
    p_ex.add_run("\n\n_________________________\nExternal Examiner\nSignature & Date").bold = True

    doc.add_page_break()

    # ==================== STUDENT DECLARATION ====================
    add_p("DECLARATION", space_before=10, space_after=24, bold=True, color=primary_color, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=16)
    decl_text = (
        f"I, {FRONT_MATTER['student_name']}, hereby declare that the project report entitled \"{FRONT_MATTER['project_title']}\", "
        f"submitted to the {FRONT_MATTER['department']}, {FRONT_MATTER['college_name']}, "
        "is an authentic record of original work carried out by me under the guidance of "
        f"{FRONT_MATTER['guide_name']}.\n\n"
        "I further declare that this work is original and that all external sources, algorithms, software libraries, and research papers "
        "have been duly cited and acknowledged. This report has not previously formed the basis for the award of any degree, diploma, "
        "associateship, or fellowship at this or any other university or institution."
    )
    add_p(decl_text, space_after=36, line_spacing=1.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    add_p("Place: Kopargaon\nDate:  ___________________", space_after=48, align=WD_ALIGN_PARAGRAPH.LEFT)
    add_p(f"_________________________\n(Signature of the Student)\n{FRONT_MATTER['student_name']}\nPRN: {FRONT_MATTER['roll_number']}", bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)

    doc.add_page_break()

    # ==================== ACKNOWLEDGMENT ====================
    add_p("ACKNOWLEDGMENT", space_before=10, space_after=20, bold=True, color=primary_color, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=16)
    ack_text = (
        "I express my deepest gratitude and sincere appreciation to my esteemed project guide, "
        f"{FRONT_MATTER['guide_name']}, for their invaluable mentorship, insightful technical guidance, and constructive "
        "critiques throughout the conceptualization, design, implementation, and evaluation of this project.\n\n"
        f"I extend my sincere thanks to Dr. Chetan D. Bawankar, Head of the Department of Cyber Security, School of Engineering and Technology (SET), "
        "Sanjivani University, Kopargaon, as well as faculty members and laboratory staff for providing excellent computational facilities and academic encouragement.\n\n"
        "Finally, I am profoundly grateful to my family, colleagues, and peers for their unwavering moral support, encouragement, "
        "and patience throughout this endeavor."
    )
    add_p(ack_text, space_after=36, line_spacing=1.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    add_p(f"{FRONT_MATTER['student_name']}\nPRN: {FRONT_MATTER['roll_number']}", bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT)

    doc.add_page_break()

    # ==================== ABSTRACT ====================
    add_p("ABSTRACT", space_before=10, space_after=20, bold=True, color=primary_color, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=16)
    for para in ABSTRACT_TEXT.split("\n\n"):
        add_p(para, space_after=10, line_spacing=1.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    doc.add_page_break()

    # ==================== TABLE OF CONTENTS ====================
    add_p("TABLE OF CONTENTS", space_before=10, space_after=18, bold=True, color=primary_color, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=16)
    toc_items = [
        ("Certificate", "iii"),
        ("Declaration", "iv"),
        ("Acknowledgment", "v"),
        ("Abstract", "vi"),
        ("List of Tables", "ix"),
        ("List of Abbreviations", "x"),
        ("Chapter 1: Introduction", "1"),
        ("  1.1 Growth of Digital Educational Content", "1"),
        ("  1.2 Intellectual Property Risks and Academic Vulnerabilities", "1"),
        ("  1.3 Motivation Behind AquaMark", "2"),
        ("  1.4 Scope and Operational Boundaries", "3"),
        ("Chapter 2: Problem Statement and Objectives", "4"),
        ("  2.1 Formal Problem Statement", "4"),
        ("  2.2 Primary Project Objectives", "4"),
        ("  2.3 Stakeholders and Intended Users", "5"),
        ("  2.4 Assumptions and Constraints", "5"),
        ("Chapter 3: Literature Review and Existing Approaches", "6"),
        ("  3.1 Survey of Document Protection Mechanisms", "6"),
        ("  3.2 Literature Comparison Table", "7"),
        ("  3.3 Identified Research and Engineering Gap", "8"),
        ("Chapter 4: Proposed System and Research Contribution", "9"),
        ("  4.1 Proposed Layered Trust Architecture", "9"),
        ("  4.2 Scientific and Engineering Contributions", "10"),
        ("Chapter 5: System Requirements and Feasibility Study", "11"),
        ("  5.1 Hardware Requirements", "11"),
        ("  5.2 Software Requirements and Runtime Environments", "11"),
        ("  5.3 Feasibility Analysis", "12"),
        ("Chapter 6: System Architecture and Detailed Design", "13"),
        ("  6.1 High-Level Architecture Overview", "13"),
        ("  6.2 Data Flow Diagrams (Level 0 and Level 1)", "13"),
        ("  6.3 Entity-Relationship (ER) Schema Design", "14"),
        ("Chapter 7: Methodology and Mathematical Algorithms", "16"),
        ("  7.1 2D-DCT Mathematical Formulation", "16"),
        ("  7.2 Steganographic Embedding Algorithm", "16"),
        ("  7.3 Steganographic Extraction Algorithm", "17"),
        ("  7.4 Cryptographic Manifest Canonicalization and Signing", "17"),
        ("Chapter 8: Implementation Details and Module Breakdown", "18"),
        ("  8.1 Authentication and Role-Based Authorization", "18"),
        ("  8.2 PDF Ingestion and Storage Engine", "18"),
        ("  8.3 Visible Watermarking Subsystem", "19"),
        ("  8.4 2D-DCT Steganographic Engine", "19"),
        ("  8.5 Multi-Path Forensic Verification Center", "20"),
        ("  8.6 Robustness Testing Laboratory", "20"),
        ("Chapter 9: Security Analysis and Threat Model", "21"),
        ("  9.1 Threat Model and Attacker Capabilities", "21"),
        ("  9.2 Threat Analysis and Mitigation Matrix", "21"),
        ("  9.3 Cryptographic Assumptions and Trust Boundaries", "22"),
        ("Chapter 10: Testing and Evaluation", "23"),
        ("  10.1 Automated Test Suite Execution and Verified Results", "23"),
        ("  10.2 Robustness Laboratory Empirical Measurements", "24"),
        ("Chapter 11: Results and Discussion", "25"),
        ("  11.1 Analysis of Experimental Findings", "25"),
        ("  11.2 Meeting Project Objectives", "26"),
        ("Chapter 12: Limitations and Future Scope", "27"),
        ("  12.1 Operational Limitations and Trade-offs", "27"),
        ("  12.2 Future Research and Engineering Enhancements", "27"),
        ("Chapter 13: Conclusion", "28"),
        ("References", "29"),
        ("Appendix A: Key Terms and Definitions", "31"),
        ("Appendix B: Core Algorithms Pseudocode", "32"),
        ("Appendix C: Test Case Execution Log", "34"),
        ("Appendix D: User Manual and Operational Walkthrough", "35"),
        ("Appendix E: Installation and Execution Guide", "36")
    ]
    for item, pg in toc_items:
        dots = ". " * int((50 - len(item)) / 1.5)
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.15
        r1 = p.add_run(item)
        r1.font.name = "Times New Roman"
        r1.font.size = Pt(10.5)
        if item.startswith("Chapter") or item.startswith("References") or item.startswith("Appendix"):
            r1.font.bold = True
        r2 = p.add_run(" " + dots + " " + pg)
        r2.font.name = "Times New Roman"
        r2.font.size = Pt(10)
        r2.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_page_break()

    # ==================== LIST OF TABLES & ABBREVIATIONS ====================
    add_p("LIST OF TABLES", space_before=10, space_after=14, bold=True, color=primary_color, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=15)
    tables_list = [
        ("Table 3.1", "Comparison Matrix of Existing Document Protection Paradigms", "7"),
        ("Table 5.1", "System Runtime Dependencies and Verified Library Versions", "12"),
        ("Table 6.1", "Database Entities and Relational Schema Specifications", "15"),
        ("Table 9.1", "Threat Model and Mitigation Matrix Under Defined Adversary Capabilities", "22"),
        ("Table 10.1", "Automated Pytest Suite Execution Specifications and Verified Outcomes", "23"),
        ("Table 10.2", "Empirical Robustness Laboratory Measurements Under Controlled Image Attacks", "24")
    ]
    for num, desc, pg in tables_list:
        add_p(f"{num}: {desc} ........................................................... {pg}", space_after=4, font_size=10.5)

    add_p("LIST OF ABBREVIATIONS", space_before=20, space_after=14, bold=True, color=primary_color, align=WD_ALIGN_PARAGRAPH.CENTER, font_size=15)
    abbrevs = [
        ("2D-DCT", "Two-Dimensional Discrete Cosine Transform"),
        ("AES-GCM", "Advanced Encryption Standard in Galois/Counter Mode"),
        ("API", "Application Programming Interface"),
        ("ASGI", "Asynchronous Server Gateway Interface"),
        ("BER", "Bit Error Rate"),
        ("DFD", "Data Flow Diagram"),
        ("DPI", "Dots Per Inch"),
        ("DRM", "Digital Rights Management"),
        ("EdDSA", "Edwards-curve Digital Signature Algorithm (Ed25519)"),
        ("HMAC", "Keyed-Hash Message Authentication Code"),
        ("HSM", "Hardware Security Module"),
        ("JWT", "JSON Web Token"),
        ("LSB", "Least Significant Bit"),
        ("ORM", "Object-Relational Mapping"),
        ("PDF", "Portable Document Format (ISO 32000-1)"),
        ("PSNR", "Peak Signal-to-Noise Ratio (in dB)"),
        ("REST", "Representational State Transfer"),
        ("SHA-256", "Secure Hash Algorithm (256-bit Digest)"),
        ("UI/UX", "User Interface / User Experience"),
        ("UUID", "Universally Unique Identifier (v4)")
    ]
    for ab, meaning in abbrevs:
        add_p(f"• {ab:<12} : {meaning}", space_after=2, font_size=10)

    doc.add_page_break()

    # ==================== CHAPTERS EXECUTION ====================
    def emit_chapter(ch_title, ch_text):
        add_h1(ch_title)
        paragraphs = ch_text.strip().split("\n\n")
        for p in paragraphs:
            lines = p.strip().split("\n")
            if len(lines) == 1 and (lines[0].startswith("1.") or lines[0].startswith("2.") or lines[0].startswith("3.") or lines[0].startswith("4.") or lines[0].startswith("5.") or lines[0].startswith("6.") or lines[0].startswith("7.") or lines[0].startswith("8.") or lines[0].startswith("9.") or lines[0].startswith("10.") or lines[0].startswith("11.") or lines[0].startswith("12.") or lines[0].startswith("13.")):
                add_h2(lines[0])
            elif lines[0].startswith("Level ") or lines[0].startswith("Summary ") or lines[0].startswith("Test Case ") or lines[0].startswith("Where "):
                add_h2(lines[0])
            elif lines[0].startswith("+--") and lines[-1].startswith("+--"):
                # Table block
                rows = [l for l in lines if not l.startswith("+--")]
                if rows:
                    h_cols = [c.strip() for c in rows[0].split("|")[1:-1]]
                    r_data = []
                    for r in rows[1:]:
                        r_data.append([c.strip() for c in r.split("|")[1:-1]])
                    add_table_grid(h_cols, r_data)
            elif lines[0].startswith("+-") or lines[0].startswith("  * ") or lines[0].startswith("1: ") or lines[0].startswith("="):
                add_code_block("\n".join(lines))
            else:
                for line in lines:
                    if line.startswith("- ") or line.startswith("• ") or line.startswith("* "):
                        add_bullet(line[2:])
                    elif line.startswith("1. ") or line.startswith("2. ") or line.startswith("3. ") or line.startswith("4. ") or line.startswith("5. ") or line.startswith("6. "):
                        add_bullet(line)
                    else:
                        add_p(line)
        doc.add_page_break()

    # Chapters
    emit_chapter(ch1_4.CHAPTER_1_TITLE, ch1_4.CHAPTER_1_TEXT)
    emit_chapter(ch1_4.CHAPTER_2_TITLE, ch1_4.CHAPTER_2_TEXT)
    emit_chapter(ch1_4.CHAPTER_3_TITLE, ch1_4.CHAPTER_3_TEXT)
    emit_chapter(ch1_4.CHAPTER_4_TITLE, ch1_4.CHAPTER_4_TEXT)
    emit_chapter(ch5_8.CHAPTER_5_TITLE, ch5_8.CHAPTER_5_TEXT)
    emit_chapter(ch5_8.CHAPTER_6_TITLE, ch5_8.CHAPTER_6_TEXT)
    emit_chapter(ch5_8.CHAPTER_7_TITLE, ch5_8.CHAPTER_7_TEXT)
    emit_chapter(ch5_8.CHAPTER_8_TITLE, ch5_8.CHAPTER_8_TEXT)
    emit_chapter(ch9_13.CHAPTER_9_TITLE, ch9_13.CHAPTER_9_TEXT)
    emit_chapter(ch9_13.CHAPTER_10_TITLE, ch9_13.CHAPTER_10_TEXT)
    emit_chapter(ch9_13.CHAPTER_11_TITLE, ch9_13.CHAPTER_11_TEXT)
    emit_chapter(ch9_13.CHAPTER_12_TITLE, ch9_13.CHAPTER_12_TEXT)
    emit_chapter(ch9_13.CHAPTER_13_TITLE, ch9_13.CHAPTER_13_TEXT)

    # REFERENCES
    add_h1("References")
    for ref_line in ch9_13.REFERENCES_TEXT.strip().split("\n"):
        add_p(ref_line, space_after=6, font_size=10)
    doc.add_page_break()

    # APPENDICES
    emit_chapter(appx.APPENDIX_A_TITLE, appx.APPENDIX_A_TEXT)
    emit_chapter(appx.APPENDIX_B_TITLE, appx.APPENDIX_B_TEXT)
    emit_chapter(appx.APPENDIX_C_TITLE, appx.APPENDIX_C_TEXT)
    emit_chapter(appx.APPENDIX_D_TITLE, appx.APPENDIX_D_TEXT)
    emit_chapter(appx.APPENDIX_E_TITLE, appx.APPENDIX_E_TEXT)

    doc.save(output_path)
    print(f"Successfully generated academic report DOCX at: {output_path}")

if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "AquaMark_Full_Project_Report.docx"
    generate_academic_report_docx(out)
